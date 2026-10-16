"""
NetBrain tables -> CSV -> Kafka.

1. Read Hostname / Tablename / VRF rows from the sheet.
2. For each row, call the NetBrain DeviceTableData API (all pages).
3. Write the response to a CSV file, with hostname and tableName columns at the end.
4. Publish one event per row to Kafka with the existing producer.
"""

import asyncio
import csv
import os
from datetime import datetime, timezone

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

BASE_URL = os.getenv("NETBRAIN_BASE_URL", "https://netbrain-prod.bank.intranet/ServicesAPI/API")
USERNAME = os.getenv("NETBRAIN_USERNAME", "")
PASSWORD = os.getenv("NETBRAIN_PASSWORD", "")
XLSX_FILE = os.getenv("INPUT_XLSX", "devices.xlsx")
OUTPUT_DIR = os.getenv("OUTPUT_DIR", "output")
PAGE_SIZE = int(os.getenv("NETBRAIN_PAGE_SIZE", "100000"))
FIRST_PAGE_INDEX = int(os.getenv("NETBRAIN_FIRST_PAGE_INDEX", "1"))
VERIFY_SSL = os.getenv("NETBRAIN_VERIFY_SSL", "true").lower() == "true"
ENTITY_TYPE = "NCT_Table"

SUCCESS = 790200
DEVICE_NOT_FOUND = 792008

# Current NetBrain token. Starts with NETBRAIN_TOKEN if given, otherwise login() fills it.
token = os.getenv("NETBRAIN_TOKEN", "")


def login():
    """Get a new token from /V1/Session."""
    global token
    if not (USERNAME and PASSWORD):
        raise Exception("Token missing or expired: set NETBRAIN_USERNAME and NETBRAIN_PASSWORD in .env")

    response = requests.post(
        f"{BASE_URL}/V1/Session",
        json={"username": USERNAME, "password": PASSWORD},
        verify=VERIFY_SSL,
        timeout=60,
    )
    response.raise_for_status()
    body = response.json()
    if body.get("statusCode") != SUCCESS or not body.get("token"):
        raise Exception(f"Login failed {body.get('statusCode')}: {body.get('statusDescription')}")
    token = body["token"]
    print("Logged in to NetBrain, got a new token")


def is_token_error(response, body):
    """True if NetBrain rejected the token (expired / invalid)."""
    if response.status_code in (401, 403):
        return True
    return "token" in str(body.get("statusDescription", "")).lower()


def get_page(params):
    """GET one page. If the token has expired, log in again and retry once."""
    url = f"{BASE_URL}/V1/CMDB/Devices/DeviceTableData"
    for attempt in range(2):
        if not token:
            login()
        headers = {"content-type": "application/json", "token": token}
        response = requests.get(url, params=params, headers=headers, verify=VERIFY_SSL, timeout=300)
        try:
            body = response.json()
        except ValueError:
            body = {}

        if attempt == 0 and is_token_error(response, body):
            print("Token expired or invalid, logging in again")
            login()
            continue

        response.raise_for_status()
        return body


def read_sheet(path):
    """Return a list of {"hostname", "table", "vrf"} from the sheet."""
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h or "").strip().lower() for h in next(rows)]

    host_col = headers.index("hoostname") if "hoostname" in headers else headers.index("hostname")
    table_col = headers.index("tablename")
    vrf_col = headers.index("vrf")

    items = []
    for row in rows:
        hostname = str(row[host_col] or "").strip()
        table = str(row[table_col] or "").strip()
        vrf = str(row[vrf_col] or "").strip()
        if hostname and table:
            items.append({"hostname": hostname, "table": table, "vrf": vrf})
    return items


def call_api(hostname, table, vrf):
    """Call the API page by page until all rows are collected.

    Returns (columns, rows), or None if NetBrain says Device Not Found.
    """
    columns, all_rows = [], []
    page_index = FIRST_PAGE_INDEX if PAGE_SIZE else 0

    while True:
        params = {"hostname": hostname, "tableName": table, "pageIndex": page_index, "pageSize": PAGE_SIZE}
        if table == "routeTable":
            params["vrf"] = vrf  # vrf only applies to route table (empty is allowed)

        body = get_page(params)
        status = body.get("statusCode")
        if status == DEVICE_NOT_FOUND:
            return None  # skip this device, caller moves on to the next hostname
        if status != SUCCESS:
            raise Exception(f"API error {status}: {body.get('statusDescription')}")

        # columns / rows can be None in the response
        columns = body.get("columns") or columns
        rows = body.get("rows") or []
        all_rows.extend(rows)

        if PAGE_SIZE == 0 or len(rows) < PAGE_SIZE:  # last page
            break
        page_index += 1

    return columns, all_rows


def write_csv(hostname, table, vrf, columns, rows):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    name = f"{hostname}_{table}" + (f"_{vrf}" if vrf else "")
    path = os.path.join(OUTPUT_DIR, "".join(c if c.isalnum() or c in "-_." else "_" for c in name) + ".csv")

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(columns + ["hostname", "tableName"])
        for row in rows:
            writer.writerow(row + [hostname, table])
    return path


def build_event(hostname, table, vrf, timestamp, columns, row):
    return {
        "eventType": table,
        "entityType": ENTITY_TYPE,
        "timestamp": timestamp,
        "attributes": {"hostname": hostname, "tableName": table, "vrf": vrf},
        "data": dict(zip(columns, row)),
    }


async def run(producer):
    items = read_sheet(XLSX_FILE)
    print(f"Read {len(items)} rows from {XLSX_FILE}")

    for i, item in enumerate(items, 1):
        hostname, table, vrf = item["hostname"], item["table"], item["vrf"]
        try:
            result = call_api(hostname, table, vrf)
            if result is None:
                print(f"[{i}/{len(items)}] {hostname} | {table} | {vrf} -> SKIPPED: Device Not Found")
                continue

            columns, rows = result
            csv_path = write_csv(hostname, table, vrf, columns, rows)

            timestamp = datetime.now(timezone.utc).isoformat()
            for row in rows:
                event_data = build_event(hostname, table, vrf, timestamp, columns, row)
                await producer.publish(key="", value=event_data)

            print(f"[{i}/{len(items)}] {hostname} | {table} | {vrf} -> {len(rows)} rows, {csv_path}")
        except Exception as e:
            print(f"[{i}/{len(items)}] {hostname} | {table} | {vrf} -> FAILED: {e}")


async def main():
    # TODO: replace with your existing producer
    from producer import producer
    await run(producer)


if __name__ == "__main__":
    asyncio.run(main())
